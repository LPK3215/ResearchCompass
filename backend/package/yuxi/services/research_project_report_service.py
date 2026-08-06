"""ResearchCompass 研究项目报告导出服务。

本模块是本仓库作者在开源智能体框架 Yuxi 之上实现的研究项目报告业务：将项目概览、
执行计划与归集成果渲染为 Markdown / DOCX 报告，供科研用户导出与归档。项目数据、
执行计划与资产可用性的获取复用 Yuxi 仓库与项目服务；报告的版式与文案由本模块决定。
"""

from __future__ import annotations

import csv
from datetime import UTC, datetime
from io import BytesIO, StringIO
from typing import Any

from docx import Document

from yuxi.repositories.research_project_repository import ResearchProjectRepository
from yuxi.services.research_project_plan_service import get_research_project_plan
from yuxi.services.research_project_service import (
    ResearchProjectError,
    get_available_project_asset_ids,
    get_owned_project,
    serialize_project,
    serialize_project_asset,
)
from yuxi.storage.postgres.models_business import User


STATUS_LABELS = {
    "active": "进行中",
    "completed": "已完成",
    "archived": "已归档",
    "planned": "待开始",
    "todo": "待处理",
    "in_progress": "进行中",
    "blocked": "受阻",
    "done": "已完成",
}
HEALTH_LABELS = {
    "completed": "已完成",
    "overdue": "已逾期",
    "at_risk": "有风险",
    "on_track": "进展正常",
    "not_planned": "尚未规划",
}
ASSET_LABELS = {
    "paper": "论文",
    "search_run": "检索记录",
    "synthesis_run": "研究综述",
    "analysis_run": "论文分析",
    "evaluation_experiment": "消融实验",
    "evidence": "证据",
}


async def export_research_project_report(
    *, project_id: str, current_user: User, export_format: str
) -> tuple[str, bytes, str]:
    if export_format not in {"markdown", "docx", "csv"}:
        raise ResearchProjectError("invalid_report_format", "报告格式只支持 markdown、docx 或 csv")
    project = await get_owned_project(project_id, current_user)
    repository = ResearchProjectRepository()
    assets = await repository.list_all_assets(project_id)
    available_asset_ids = await get_available_project_asset_ids(project, current_user, assets, repository=repository)
    counts = {asset_type: 0 for asset_type in ASSET_LABELS}
    for asset in assets:
        counts[asset.asset_type] += 1
    counts["total"] = len(assets)
    plan = await get_research_project_plan(
        project_id=project_id,
        current_user=current_user,
        project=project,
        available_asset_ids=available_asset_ids,
    )
    project_payload = serialize_project(project, counts, plan["summary"])
    asset_payloads = [
        serialize_project_asset(asset, available=asset.asset_id in available_asset_ids) for asset in assets
    ]
    report = {
        "project": project_payload,
        "plan": plan,
        "assets": asset_payloads,
        "generated_at": datetime.now(UTC).isoformat(),
    }
    filename_base = f"research-project-{project.project_id}"
    if export_format == "markdown":
        return f"{filename_base}.md", render_project_markdown(report).encode("utf-8"), "text/markdown; charset=utf-8"
    if export_format == "csv":
        return f"{filename_base}.csv", render_project_csv(report), "text/csv; charset=utf-8"
    return (
        f"{filename_base}.docx",
        render_project_docx(report),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def render_project_markdown(report: dict[str, Any]) -> str:
    project = report["project"]
    project_id = str(project.get("project_id") or "project")
    plan = report["plan"]
    progress_source = "任务自动计算" if project["progress_source"] == "tasks" else "手动维护"
    lines = [
        f"# {project['title']}",
        "",
        "## 项目概览",
        "",
        f"- 研究问题：{project['research_question']}",
        f"- 项目状态：{STATUS_LABELS.get(project['status'], project['status'])}",
        f"- 项目健康度：{HEALTH_LABELS.get(project['health'], project['health'])}",
        f"- 当前进度：{project['progress']}%（{progress_source}）",
        f"- 目标日期：{project['target_date'] or '未设置'}",
        f"- 下一截止日期：{project['next_due_date'] or '未设置'}",
        f"- 下一步行动：{project['next_action'] or '未设置'}",
        "",
        project["description"] or "未填写项目说明。",
        "",
        "## 执行计划",
        "",
    ]
    if not plan["milestones"] and not plan["unassigned_tasks"]:
        lines.extend(["尚未创建里程碑或任务。", ""])
    for index, milestone in enumerate(plan["milestones"], start=1):
        lines.extend(
            [
                f"### {index}. {milestone['title']}",
                "",
                (
                    f"状态：{STATUS_LABELS.get(milestone['status'], milestone['status'])} | "
                    f"进度：{milestone['progress']}% | 目标日期：{milestone['target_date'] or '未设置'}"
                ),
                "",
            ]
        )
        if milestone["description"]:
            lines.extend([milestone["description"], ""])
        _append_markdown_tasks(lines, milestone["tasks"])
        _append_markdown_links(lines, milestone["asset_links"], "里程碑证据")
    if plan["unassigned_tasks"]:
        lines.extend(["### 未分配任务", ""])
        _append_markdown_tasks(lines, plan["unassigned_tasks"])
    lines.extend(["## 全部研究成果", ""])
    if not report["assets"]:
        lines.extend(["尚未归集研究成果。", ""])
    else:
        for asset in report["assets"]:
            availability = "可访问" if asset["available"] else "源成果已失效，仅保留项目快照"
            lines.append(
                f"- **{asset['title']}** · "
                f"{ASSET_LABELS.get(asset['asset_type'], asset['asset_type'])} · {availability}"
            )
            if asset["notes"]:
                lines.append(f"  - 项目备注：{asset['notes']}")
        lines.append("")
    lines.extend(["---", "", f"报告生成时间：{report['generated_at']}", ""])
    return "\n".join(lines)


def render_project_csv(report: dict[str, Any]) -> bytes:
    """Render a flat, audit-friendly project snapshot for spreadsheet workflows."""
    output = StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    _write_csv_row(writer, [
            "record_type",
            "record_id",
            "parent_id",
            "title",
            "status",
            "priority",
            "due_date",
            "asset_type",
            "available",
            "notes",
    ])
    project = report["project"]
    project_id = str(project.get("project_id") or "project")
    _write_csv_row(writer, [
            "project",
            project_id,
            "",
            project.get("title", ""),
            project.get("status", ""),
            "",
            project.get("target_date", "") or "",
            "",
            "true",
            project.get("description", "") or "",
    ])
    for milestone in report["plan"].get("milestones", []):
        milestone_id = milestone.get("milestone_id") or f"milestone-{milestone.get('title', 'unknown')}"
        _write_csv_row(writer, [
            "milestone",
            milestone_id,
            project_id,
                milestone.get("title", ""),
                milestone.get("status", ""),
                "",
                milestone.get("target_date", "") or "",
                "",
                "true",
                milestone.get("description", "") or "",
        ])
        for task in milestone.get("tasks", []):
            _write_csv_row(writer, _task_csv_row(task, milestone_id))
    for task in report["plan"].get("unassigned_tasks", []):
        _write_csv_row(writer, _task_csv_row(task, project_id))
    for asset in report.get("assets", []):
        _write_csv_row(writer, [
                "asset",
                asset.get("asset_id", ""),
                project_id,
                asset.get("title", ""),
                "",
                "",
                "",
                asset.get("asset_type", ""),
                str(bool(asset.get("available"))).lower(),
                asset.get("notes", "") or "",
        ])
    return output.getvalue().encode("utf-8-sig")


def _task_csv_row(task: dict[str, Any], parent_id: str) -> list[str]:
    return [
        "task",
        task.get("task_id", ""),
        parent_id,
        task.get("title", ""),
        task.get("status", ""),
        task.get("priority", ""),
        task.get("due_date", "") or "",
        "",
        "true",
        task.get("description", "") or "",
    ]


def _csv_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _write_csv_row(writer: csv.writer, row: list[Any]) -> None:
    writer.writerow([_csv_cell(value) for value in row])


def _append_markdown_tasks(lines: list[str], tasks: list[dict[str, Any]]) -> None:
    if not tasks:
        lines.extend(["暂无任务。", ""])
        return
    for task in tasks:
        marker = "x" if task["status"] == "done" else " "
        details = (
            f"{STATUS_LABELS.get(task['status'], task['status'])} / {task['priority']} / "
            f"截止 {task['due_date'] or '未设置'}"
        )
        lines.append(f"- [{marker}] **{task['title']}**（{details}）")
        if task["description"]:
            lines.append(f"  - {task['description']}")
        for link in task["asset_links"]:
            asset = link["asset"]
            availability = "" if asset["available"] else "（源成果已失效）"
            lines.append(f"  - 证据：{asset['title']}{availability}")
    lines.append("")


def _append_markdown_links(lines: list[str], links: list[dict[str, Any]], heading: str) -> None:
    if not links:
        return
    lines.append(f"**{heading}**")
    for link in links:
        asset = link["asset"]
        availability = "" if asset["available"] else "（源成果已失效）"
        lines.append(f"- {asset['title']}{availability}")
    lines.append("")


def render_project_docx(report: dict[str, Any]) -> bytes:
    project = report["project"]
    plan = report["plan"]
    document = Document()
    document.add_heading(project["title"], level=0)
    document.add_heading("项目概览", level=1)
    overview = [
        ("研究问题", project["research_question"]),
        ("项目状态", STATUS_LABELS.get(project["status"], project["status"])),
        ("项目健康度", HEALTH_LABELS.get(project["health"], project["health"])),
        ("当前进度", f"{project['progress']}%"),
        ("目标日期", project["target_date"] or "未设置"),
        ("下一截止日期", project["next_due_date"] or "未设置"),
        ("下一步行动", project["next_action"] or "未设置"),
    ]
    for label, value in overview:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(str(value))
    if project["description"]:
        document.add_paragraph(project["description"])

    document.add_heading("执行计划", level=1)
    if not plan["milestones"] and not plan["unassigned_tasks"]:
        document.add_paragraph("尚未创建里程碑或任务。")
    for index, milestone in enumerate(plan["milestones"], start=1):
        document.add_heading(f"{index}. {milestone['title']}", level=2)
        document.add_paragraph(
            f"状态：{STATUS_LABELS.get(milestone['status'], milestone['status'])} | "
            f"进度：{milestone['progress']}% | 目标日期：{milestone['target_date'] or '未设置'}"
        )
        if milestone["description"]:
            document.add_paragraph(milestone["description"])
        _append_docx_tasks(document, milestone["tasks"])
        for link in milestone["asset_links"]:
            asset = link["asset"]
            suffix = "（源成果已失效）" if not asset["available"] else ""
            document.add_paragraph(f"里程碑证据：{asset['title']}{suffix}", style="List Bullet")
    if plan["unassigned_tasks"]:
        document.add_heading("未分配任务", level=2)
        _append_docx_tasks(document, plan["unassigned_tasks"])

    document.add_heading("全部研究成果", level=1)
    if not report["assets"]:
        document.add_paragraph("尚未归集研究成果。")
    for asset in report["assets"]:
        availability = "可访问" if asset["available"] else "源成果已失效，仅保留项目快照"
        document.add_paragraph(
            f"{asset['title']} · {ASSET_LABELS.get(asset['asset_type'], asset['asset_type'])} · {availability}",
            style="List Bullet",
        )
        if asset["notes"]:
            document.add_paragraph(f"项目备注：{asset['notes']}")
    document.add_paragraph(f"报告生成时间：{report['generated_at']}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _append_docx_tasks(document: Document, tasks: list[dict[str, Any]]) -> None:
    if not tasks:
        document.add_paragraph("暂无任务。")
        return
    for task in tasks:
        marker = "[x]" if task["status"] == "done" else "[ ]"
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{marker} {task['title']}").bold = True
        paragraph.add_run(
            f"（{STATUS_LABELS.get(task['status'], task['status'])} / {task['priority']} / "
            f"截止 {task['due_date'] or '未设置'}）"
        )
        if task["description"]:
            document.add_paragraph(task["description"])
        for link in task["asset_links"]:
            asset = link["asset"]
            suffix = "（源成果已失效）" if not asset["available"] else ""
            document.add_paragraph(f"证据：{asset['title']}{suffix}", style="List Bullet 2")


__all__ = ["export_research_project_report", "render_project_csv", "render_project_docx", "render_project_markdown"]
