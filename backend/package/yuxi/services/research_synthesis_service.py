"""ResearchCompass 证据约束研究综述服务。

本模块是本仓库作者在开源智能体框架 Yuxi 之上设计的跨论文综述业务：以本地混合检索
快照为唯一证据来源，调用大模型生成结构化综述，再对每条结论的证据 chunk_id 做身份
复核与引用覆盖率校验，确保结论只引用真实存在的检索证据。模型调用、任务调度与检索
运行时由 Yuxi 提供；本模块定义证据约束的提示词契约、可验证综述的校验规则与任务
可恢复执行协议。
"""

from __future__ import annotations

import asyncio
import json
import re
import time
import uuid
from collections.abc import Mapping
from io import BytesIO
from typing import Any

from docx import Document
from fastapi import HTTPException
from sqlalchemy.exc import IntegrityError

from yuxi.config import config
from yuxi.knowledge.chunking.ragflow_like.nlp import count_tokens
from yuxi.models import select_model
from yuxi.models.providers.cache import model_cache
from yuxi.repositories.academic_paper_repository import AcademicPaperRepository
from yuxi.repositories.knowledge_base_repository import KnowledgeBaseRepository
from yuxi.repositories.knowledge_chunk_repository import KnowledgeChunkRepository
from yuxi.repositories.research_synthesis_repository import ResearchSynthesisRepository
from yuxi.repositories.user_repository import UserRepository
from yuxi.services.research_paper_service import _ensure_access
from yuxi.services.research_search_service import LOCAL_HYBRID_MODE, ResearchSearchError, search_papers
from yuxi.services.task_service import PublicTaskError, TaskContext, tasker
from yuxi.storage.postgres.models_business import User
from yuxi.utils.datetime_utils import utc_now_naive


class ResearchSynthesisError(PublicTaskError):
    def __init__(self, error_type: str, message: str):
        super().__init__(message)
        self.error_type = error_type
        self.message = message


DEFAULT_SYNTHESIS_INPUT_BUDGET = 24_000
SYNTHESIS_CONTEXT_OVERHEAD = 1_024
SYNTHESIS_CONTEXT_RATIO = 0.75
SYNTHESIS_RULES_VERSION = "evidence-synthesis-v1"
CONFIDENCE_LEVELS = {"high", "medium", "low"}

SYNTHESIS_SYSTEM_PROMPT = """你是证据约束的跨论文研究综述专家。只能使用输入中的论文和证据 chunk。
禁止引入外部事实，禁止编造 chunk_id 或 paper_id。
只返回合法 JSON 对象，不要 Markdown，不要额外解释，严格使用以下结构：
{
  "executive_summary":{"text":"...","claim_ids":["C1"]},
  "themes":[{"title":"...","summary":"...","claim_ids":["C1"]}],
  "claims":[{"claim_id":"C1","statement":"...","confidence":"high|medium|low","evidence_chunk_ids":["chunk-id"]}],
  "contradictions":[{"statement":"...","evidence_chunk_ids":["chunk-a","chunk-b"]}],
  "limitations":[{"statement":"...","evidence_chunk_ids":["chunk-id"]}],
  "research_gaps":[{"statement":"...","basis":"...","confidence":"high|medium|low","evidence_chunk_ids":["chunk-id"]}],
  "unsupported_claims":[{"statement":"...","reason":"..."}]
}
规则：每个受支持结论必须引用输入中真实存在的 evidence_chunk_ids。
executive_summary 和 themes 只能概括其 claim_ids 对应的 claims。
high 置信度至少需要两篇不同论文的证据；contradictions 至少需要两篇不同论文。
证据不足的判断放入 unsupported_claims，不能伪装成结论。"""


def _elapsed(started: float) -> int:
    return round((time.perf_counter() - started) * 1000)


def _parse_json(content: str) -> dict[str, Any]:
    value = re.sub(r"^```(?:json)?\s*|\s*```$", "", content.strip(), flags=re.IGNORECASE | re.DOTALL).strip()
    try:
        payload = json.loads(value)
    except json.JSONDecodeError as exc:
        raise ResearchSynthesisError("synthesis_invalid_json", "综述模型返回的 JSON 不可解析") from exc
    if not isinstance(payload, dict):
        raise ResearchSynthesisError("synthesis_invalid_schema", "综述模型必须返回 JSON 对象")
    return payload


def _model_input_budget(model: Any) -> tuple[int, str]:
    candidates: list[tuple[int, str]] = []
    profile = getattr(getattr(model, "model", None), "profile", None)
    if isinstance(profile, Mapping):
        max_input_tokens = profile.get("max_input_tokens")
        if isinstance(max_input_tokens, int) and max_input_tokens > 0:
            candidates.append((max_input_tokens, "模型运行时输入上限"))

    model_info = getattr(model, "info", {})
    context_length = model_info.get("context_length") if isinstance(model_info, Mapping) else None
    if isinstance(context_length, int) and context_length > 0:
        candidates.append(
            (max(int(context_length * SYNTHESIS_CONTEXT_RATIO), 256), "模型配置 context_length 的 75% 安全预算")
        )
    if not candidates:
        return DEFAULT_SYNTHESIS_INPUT_BUDGET, "未声明模型容量时的默认安全预算"
    return min(candidates, key=lambda item: item[0])


def _ensure_context_budget(model: Any, context: str) -> None:
    estimated_tokens = count_tokens(SYNTHESIS_SYSTEM_PROMPT) + count_tokens(context) + SYNTHESIS_CONTEXT_OVERHEAD
    budget, budget_source = _model_input_budget(model)
    if estimated_tokens <= budget:
        return
    raise ResearchSynthesisError(
        "synthesis_context_exceeded",
        f"综述预计输入约 {estimated_tokens:,} tokens，超过当前模型的 {budget:,} tokens 安全预算（{budget_source}）。"
        "系统不会静默丢弃检索证据；请减少返回论文数，或选择更大上下文模型。",
    )


def _required_text(value: Any, field: str, *, max_length: int = 12_000) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ResearchSynthesisError("synthesis_invalid_schema", f"{field} 必须是非空字符串")
    normalized = value.strip()
    if len(normalized) > max_length:
        raise ResearchSynthesisError("synthesis_invalid_schema", f"{field} 超过允许长度")
    return normalized


def _required_list(value: Any, field: str, *, max_items: int) -> list[Any]:
    if not isinstance(value, list) or len(value) > max_items:
        raise ResearchSynthesisError("synthesis_invalid_schema", f"{field} 必须是最多 {max_items} 项的数组")
    return value


def _string_ids(value: Any, field: str, *, max_items: int = 100) -> list[str]:
    items = _required_list(value, field, max_items=max_items)
    normalized = [item.strip() for item in items if isinstance(item, str) and item.strip()]
    if len(normalized) != len(items):
        raise ResearchSynthesisError("synthesis_invalid_schema", f"{field} 只能包含非空字符串")
    return list(dict.fromkeys(normalized))


def _build_snapshot(search_result: dict[str, Any]) -> dict[str, Any]:
    papers: list[dict[str, Any]] = []
    seen_chunks: set[str] = set()
    for item in search_result.get("items") or []:
        paper_id = _required_text(item.get("paper_id"), "retrieval.paper_id", max_length=64)
        evidence_rows = []
        for evidence in item.get("evidence") or []:
            chunk_id = _required_text(evidence.get("chunk_id"), "retrieval.chunk_id", max_length=128)
            file_id = _required_text(evidence.get("file_id"), "retrieval.file_id", max_length=64)
            content = _required_text(evidence.get("content"), "retrieval.content", max_length=200_000)
            if chunk_id in seen_chunks:
                continue
            seen_chunks.add(chunk_id)
            evidence_rows.append(
                {
                    "chunk_id": chunk_id,
                    "file_id": file_id,
                    "content": content,
                    "section_type": evidence.get("section_type"),
                    "section_title": evidence.get("section_title"),
                    "locator": evidence.get("locator") or {},
                }
            )
        if evidence_rows:
            papers.append(
                {
                    "paper_id": paper_id,
                    "title": item.get("title") or "未命名论文",
                    "authors": item.get("authors") or [],
                    "publication_year": item.get("publication_year"),
                    "venue": item.get("venue"),
                    "doi": item.get("doi"),
                    "ranking_score": item.get("ranking_score"),
                    "evidence": evidence_rows,
                }
            )
    return {
        "search_run_id": search_result.get("run_id"),
        "rewritten_query": search_result.get("rewritten_query"),
        "keywords": search_result.get("keywords") or [],
        "config": search_result.get("config") or {},
        "papers": papers,
    }


def _model_context(query: str, snapshot: dict[str, Any]) -> str:
    payload = {
        "research_question": query,
        "retrieval": {
            "search_run_id": snapshot.get("search_run_id"),
            "rewritten_query": snapshot.get("rewritten_query"),
            "keywords": snapshot.get("keywords") or [],
        },
        "papers": snapshot.get("papers") or [],
    }
    return json.dumps(payload, ensure_ascii=False)


async def _call_synthesis_model(model_spec: str, context: str) -> dict[str, Any]:
    model = select_model(model_spec=model_spec, model_params={"temperature": 0})
    _ensure_context_budget(model, context)
    try:
        response = await model.call(
            [
                {"role": "system", "content": SYNTHESIS_SYSTEM_PROMPT},
                {"role": "user", "content": context},
            ],
            stream=False,
        )
    except ResearchSynthesisError:
        raise
    except Exception as exc:
        raise ResearchSynthesisError("synthesis_generation_failed", "综述模型调用失败") from exc
    return _parse_json(str(response.content or ""))


async def _refresh_synthesis_owner(current_user: User, kb_id: str) -> User:
    uid = str(getattr(current_user, "uid", "") or "").strip()
    if not uid:
        raise ResearchSynthesisError("forbidden", "研究综述任务缺少有效所有者")
    user = await UserRepository().get_by_uid(uid)
    if user is None or bool(user.is_deleted):
        raise ResearchSynthesisError("forbidden", "综述任务所有者不存在或已被删除")
    try:
        await _ensure_access(user, kb_id)
    except HTTPException as exc:
        raise ResearchSynthesisError("forbidden", "综述任务所有者已失去知识库访问权限") from exc
    return user


async def _verified_evidence_index(kb_id: str, snapshot: dict[str, Any]) -> dict[str, dict[str, Any]]:
    snapshot_rows: dict[str, tuple[str, dict[str, Any], dict[str, Any]]] = {}
    for paper in snapshot.get("papers") or []:
        for evidence in paper.get("evidence") or []:
            snapshot_rows[str(evidence["chunk_id"])] = (str(paper["paper_id"]), paper, evidence)

    chunk_ids = list(snapshot_rows)
    chunks = await KnowledgeChunkRepository().list_by_chunk_ids(chunk_ids)
    chunks_by_id = {str(chunk.chunk_id): chunk for chunk in chunks}
    if set(chunks_by_id) != set(chunk_ids):
        raise ResearchSynthesisError("synthesis_evidence_changed", "部分检索证据已被删除，不能生成可验证综述")

    file_ids = list(dict.fromkeys(str(chunk.file_id) for chunk in chunks))
    papers = await AcademicPaperRepository().list_by_file_ids(kb_id=kb_id, file_ids=file_ids)
    papers_by_file = {str(paper.file_id): paper for paper in papers}
    verified: dict[str, dict[str, Any]] = {}
    for chunk_id, (snapshot_paper_id, snapshot_paper, evidence) in snapshot_rows.items():
        chunk = chunks_by_id[chunk_id]
        paper = papers_by_file.get(str(chunk.file_id))
        if (
            str(chunk.kb_id) != kb_id
            or str(chunk.file_id) != str(evidence["file_id"])
            or str(chunk.content) != str(evidence["content"])
            or paper is None
            or str(paper.paper_id) != snapshot_paper_id
        ):
            raise ResearchSynthesisError("synthesis_evidence_changed", "检索证据与当前知识库论文身份不一致")
        verified[chunk_id] = {
            "chunk_id": chunk_id,
            "paper_id": str(paper.paper_id),
            "paper_title": str(paper.title),
            "section_type": evidence.get("section_type"),
            "section_title": evidence.get("section_title"),
            "locator": evidence.get("locator") or {},
        }
    return verified


def _citations(chunk_ids: list[str], evidence_index: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    unknown = [chunk_id for chunk_id in chunk_ids if chunk_id not in evidence_index]
    if unknown:
        raise ResearchSynthesisError(
            "synthesis_invalid_citation",
            f"综述模型引用了检索快照之外的证据 chunk: {', '.join(unknown[:5])}",
        )
    return [evidence_index[chunk_id] for chunk_id in chunk_ids]


def _validate_report(
    payload: dict[str, Any],
    *,
    query: str,
    snapshot: dict[str, Any],
    evidence_index: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    unsupported: list[dict[str, str]] = []
    adjustments: list[dict[str, str]] = []

    for item in _required_list(payload.get("unsupported_claims"), "unsupported_claims", max_items=60):
        if not isinstance(item, dict):
            raise ResearchSynthesisError("synthesis_invalid_schema", "unsupported_claims 的每一项必须是对象")
        unsupported.append(
            {
                "statement": _required_text(item.get("statement"), "unsupported_claims.statement", max_length=4_000),
                "reason": _required_text(item.get("reason"), "unsupported_claims.reason", max_length=2_000),
            }
        )

    claims: list[dict[str, Any]] = []
    claim_ids: set[str] = set()
    for item in _required_list(payload.get("claims"), "claims", max_items=60):
        if not isinstance(item, dict):
            raise ResearchSynthesisError("synthesis_invalid_schema", "claims 的每一项必须是对象")
        claim_id = _required_text(item.get("claim_id"), "claims.claim_id", max_length=64)
        if claim_id in claim_ids:
            raise ResearchSynthesisError("synthesis_invalid_schema", f"claim_id 重复: {claim_id}")
        claim_ids.add(claim_id)
        statement = _required_text(item.get("statement"), "claims.statement", max_length=4_000)
        chunk_ids = _string_ids(item.get("evidence_chunk_ids"), "claims.evidence_chunk_ids")
        if not chunk_ids:
            unsupported.append({"statement": statement, "reason": "模型未提供可验证证据"})
            continue
        evidence = _citations(chunk_ids, evidence_index)
        paper_ids = list(dict.fromkeys(citation["paper_id"] for citation in evidence))
        confidence = str(item.get("confidence") or "").strip().lower()
        if confidence not in CONFIDENCE_LEVELS:
            raise ResearchSynthesisError("synthesis_invalid_schema", "claims.confidence 必须是 high、medium 或 low")
        if confidence == "high" and len(paper_ids) < 2:
            confidence = "medium"
            adjustments.append({"claim_id": claim_id, "reason": "高置信结论只有一篇论文支持，已按验证规则降为 medium"})
        claims.append(
            {
                "claim_id": claim_id,
                "statement": statement,
                "confidence": confidence,
                "paper_ids": paper_ids,
                "evidence": evidence,
            }
        )

    supported_claim_ids = {claim["claim_id"] for claim in claims}
    executive = payload.get("executive_summary")
    if not isinstance(executive, dict):
        raise ResearchSynthesisError("synthesis_invalid_schema", "executive_summary 必须是对象")
    summary_claim_ids = _string_ids(executive.get("claim_ids"), "executive_summary.claim_ids", max_items=60)
    if any(claim_id not in supported_claim_ids for claim_id in summary_claim_ids):
        raise ResearchSynthesisError("synthesis_invalid_schema", "executive_summary 引用了未通过验证的 claim_id")
    summary_text = _required_text(executive.get("text"), "executive_summary.text")
    if claims and not summary_claim_ids:
        raise ResearchSynthesisError("synthesis_invalid_schema", "executive_summary 必须引用至少一个已验证 claim_id")
    if not claims:
        summary_text = "检索证据未形成通过验证的核心结论，请查看未支持结论与方法边界。"
        summary_claim_ids = []

    themes: list[dict[str, Any]] = []
    for item in _required_list(payload.get("themes"), "themes", max_items=12):
        if not isinstance(item, dict):
            raise ResearchSynthesisError("synthesis_invalid_schema", "themes 的每一项必须是对象")
        theme_claim_ids = _string_ids(item.get("claim_ids"), "themes.claim_ids", max_items=60)
        if not theme_claim_ids or any(claim_id not in supported_claim_ids for claim_id in theme_claim_ids):
            raise ResearchSynthesisError("synthesis_invalid_schema", "themes 必须且只能引用已验证 claim_id")
        themes.append(
            {
                "title": _required_text(item.get("title"), "themes.title", max_length=500),
                "summary": _required_text(item.get("summary"), "themes.summary", max_length=4_000),
                "claim_ids": theme_claim_ids,
            }
        )

    def validate_evidence_statements(field: str, max_items: int, *, require_two_papers: bool = False):
        validated = []
        for item in _required_list(payload.get(field), field, max_items=max_items):
            if not isinstance(item, dict):
                raise ResearchSynthesisError("synthesis_invalid_schema", f"{field} 的每一项必须是对象")
            statement = _required_text(item.get("statement"), f"{field}.statement", max_length=4_000)
            chunk_ids = _string_ids(item.get("evidence_chunk_ids"), f"{field}.evidence_chunk_ids")
            if not chunk_ids:
                unsupported.append({"statement": statement, "reason": f"{field} 未提供可验证证据"})
                continue
            evidence = _citations(chunk_ids, evidence_index)
            paper_ids = list(dict.fromkeys(citation["paper_id"] for citation in evidence))
            if require_two_papers and len(paper_ids) < 2:
                unsupported.append({"statement": statement, "reason": "矛盾判断不足两篇不同论文证据"})
                continue
            row: dict[str, Any] = {"statement": statement, "paper_ids": paper_ids, "evidence": evidence}
            if field == "research_gaps":
                row["basis"] = _required_text(item.get("basis"), "research_gaps.basis", max_length=4_000)
                confidence = str(item.get("confidence") or "").strip().lower()
                if confidence not in CONFIDENCE_LEVELS:
                    raise ResearchSynthesisError(
                        "synthesis_invalid_schema", "research_gaps.confidence 必须是 high、medium 或 low"
                    )
                if confidence == "high" and len(paper_ids) < 2:
                    confidence = "medium"
                    adjustments.append(
                        {"claim_id": f"research_gap_{len(validated) + 1}", "reason": "高置信研究空白只有一篇论文支持"}
                    )
                row["confidence"] = confidence
            validated.append(row)
        return validated

    contradictions = validate_evidence_statements("contradictions", 20, require_two_papers=True)
    limitations = validate_evidence_statements("limitations", 30)
    research_gaps = validate_evidence_statements("research_gaps", 30)
    supported_conclusions = len(claims) + len(contradictions) + len(limitations) + len(research_gaps)
    total_conclusions = supported_conclusions + len(unsupported)
    cited_paper_ids = {
        paper_id
        for collection in (claims, contradictions, limitations, research_gaps)
        for item in collection
        for paper_id in item["paper_ids"]
    }
    snapshot_papers = snapshot.get("papers") or []
    snapshot_chunk_count = sum(len(paper.get("evidence") or []) for paper in snapshot_papers)
    coverage = {
        "supported_claims": len(claims),
        "total_claims": len(claims) + len(unsupported),
        "supported_conclusions": supported_conclusions,
        "total_conclusions": total_conclusions,
        "citation_coverage_ratio": round(supported_conclusions / total_conclusions, 4) if total_conclusions else 0.0,
        "distinct_papers": len(cited_paper_ids),
        "retrieved_papers": len(snapshot_papers),
        "retrieved_chunks": snapshot_chunk_count,
    }
    return {
        "question": query,
        "executive_summary": {"text": summary_text, "claim_ids": summary_claim_ids},
        "themes": themes,
        "claims": claims,
        "contradictions": contradictions,
        "limitations": limitations,
        "research_gaps": research_gaps,
        "unsupported_claims": unsupported,
        "coverage": coverage,
        "sources": [
            {
                "paper_id": paper["paper_id"],
                "title": paper["title"],
                "authors": paper.get("authors") or [],
                "publication_year": paper.get("publication_year"),
                "venue": paper.get("venue"),
                "doi": paper.get("doi"),
                "evidence_chunk_ids": [row["chunk_id"] for row in paper.get("evidence") or []],
            }
            for paper in snapshot_papers
        ],
        "retrieval": {
            "mode": LOCAL_HYBRID_MODE,
            "search_run_id": snapshot.get("search_run_id"),
            "rewritten_query": snapshot.get("rewritten_query"),
            "keywords": snapshot.get("keywords") or [],
        },
        "methodological_constraints": [
            "报告仅基于当前知识库的本地混合检索快照，不代表外部学术数据库的完整覆盖。",
            "引用覆盖率衡量结论是否绑定检索证据，不等同于结论在整个学术领域中的绝对正确性。",
        ],
        "validation": {
            "status": "verified",
            "rules_version": SYNTHESIS_RULES_VERSION,
            "validated_at": utc_now_naive().isoformat(),
            "confidence_adjustments": adjustments,
        },
    }


async def _update_run_or_cancel(
    repo: ResearchSynthesisRepository,
    context: TaskContext,
    run_id: str,
    values: dict[str, Any],
) -> None:
    updated = await repo.update_if_not_cancelled(run_id, values)
    if updated is not None:
        return
    context.cancellation_reason = "cancelled"
    raise asyncio.CancelledError("Synthesis was cancelled before a state transition")


async def _run_synthesis(
    context: TaskContext,
    *,
    run_id: str,
    kb_id: str,
    current_user: User,
    query: str,
    model_spec: str,
    reranker_model: str,
    retrieval_config: dict[str, Any],
    persisted_snapshot: dict[str, Any] | None = None,
    resume_stage_timings: dict[str, Any] | None = None,
) -> dict[str, Any]:
    repo = ResearchSynthesisRepository()
    timings: dict[str, Any] = dict(resume_stage_timings or {})
    snapshot = (
        persisted_snapshot
        if isinstance(persisted_snapshot, dict)
        and isinstance(persisted_snapshot.get("papers"), list)
        and len(persisted_snapshot["papers"]) >= 2
        else None
    )
    if snapshot is None:
        await _update_run_or_cancel(
            repo,
            context,
            run_id,
            {
                "status": "retrieving",
                "stage": "retrieving",
                "started_at": utc_now_naive(),
                "completed_at": None,
                "error_type": None,
                "error_message": None,
                "stage_timings": timings,
            },
        )
    else:
        await _update_run_or_cancel(
            repo,
            context,
            run_id,
            {
                "status": "synthesizing",
                "stage": "synthesizing",
                "retrieval_snapshot": snapshot,
                "stage_timings": timings,
            },
        )
    try:
        await context.raise_if_cancelled()
        current_user = await _refresh_synthesis_owner(current_user, kb_id)
        if snapshot is None:
            await context.set_progress(5, "正在执行本地混合检索")
            started = time.perf_counter()
            try:
                search_result = await search_papers(
                    kb_id=kb_id,
                    current_user=current_user,
                    query=query,
                    retrieval_mode=LOCAL_HYBRID_MODE,
                    top_k=int(retrieval_config["top_k"]),
                    recall_top_k=int(retrieval_config["recall_top_k"]),
                    year_from=retrieval_config.get("year_from"),
                    year_to=retrieval_config.get("year_to"),
                    chat_model=model_spec,
                    reranker_model=reranker_model,
                )
            except ResearchSearchError as exc:
                raise ResearchSynthesisError("synthesis_retrieval_failed", exc.message) from exc
            timings["retrieval_ms"] = _elapsed(started)
            snapshot = _build_snapshot(search_result)
            if len(snapshot["papers"]) < 2:
                raise ResearchSynthesisError(
                    "synthesis_insufficient_evidence",
                    "跨论文综述至少需要检索命中两篇含证据的论文",
                )
            await _update_run_or_cancel(
                repo,
                context,
                run_id,
                {
                    "status": "synthesizing",
                    "stage": "synthesizing",
                    "retrieval_snapshot": snapshot,
                    "stage_timings": timings,
                },
            )
        else:
            await context.set_progress(40, "服务重启，复用已持久化的检索快照")

        await context.raise_if_cancelled()
        current_user = await _refresh_synthesis_owner(current_user, kb_id)
        await context.set_progress(45, "正在生成结构化跨论文综述")
        started = time.perf_counter()
        raw_report = await _call_synthesis_model(model_spec, _model_context(query, snapshot))
        timings["synthesis_ms"] = _elapsed(started)
        await _update_run_or_cancel(
            repo,
            context,
            run_id,
            {"status": "validating", "stage": "validating", "stage_timings": timings},
        )

        await context.raise_if_cancelled()
        current_user = await _refresh_synthesis_owner(current_user, kb_id)
        await context.set_progress(80, "正在校验证据身份与引用覆盖率")
        started = time.perf_counter()
        evidence_index = await _verified_evidence_index(kb_id, snapshot)
        result = _validate_report(raw_report, query=query, snapshot=snapshot, evidence_index=evidence_index)
        timings["validation_ms"] = _elapsed(started)
        await context.raise_if_cancelled()
        await _refresh_synthesis_owner(current_user, kb_id)
        completed = await repo.update_if_not_cancelled(
            run_id,
            {
                "status": "success",
                "stage": "completed",
                "result": result,
                "stage_timings": timings,
                "completed_at": utc_now_naive(),
            },
        )
        if completed is None:
            context.cancellation_reason = "cancelled"
            raise asyncio.CancelledError("Synthesis was cancelled before completion")
        await context.set_result(result)
        await context.set_progress(100, "证据约束研究综述生成完成")
        return result
    except asyncio.CancelledError:
        if context.cancellation_reason == "shutdown":
            await repo.update_if_not_cancelled(
                run_id,
                {
                    "status": "pending",
                    "stage": "pending",
                    "stage_timings": timings,
                    "error_type": "synthesis_recovery_pending",
                    "error_message": "服务重启，综述任务等待自动恢复",
                    "completed_at": None,
                },
            )
        else:
            timed_out = context.cancellation_reason == "timeout"
            await repo.update_if_not_cancelled(
                run_id,
                {
                    "status": "failed" if timed_out else "cancelled",
                    "stage": "failed" if timed_out else "cancelled",
                    "stage_timings": timings,
                    "error_type": "synthesis_timeout" if timed_out else "synthesis_cancelled",
                    "error_message": "综述任务执行超时" if timed_out else "综述任务已取消",
                    "completed_at": utc_now_naive(),
                },
            )
        raise
    except Exception as exc:
        if isinstance(exc, (ResearchSynthesisError, ResearchSearchError)):
            failure = ResearchSynthesisError(exc.error_type, exc.message)
        else:
            failure = ResearchSynthesisError("synthesis_failed", "综述任务执行失败")
        await repo.update_if_not_cancelled(
            run_id,
            {
                "status": "failed",
                "stage": "failed",
                "stage_timings": timings,
                "error_type": failure.error_type,
                "error_message": failure.message,
                "completed_at": utc_now_naive(),
            },
        )
        if isinstance(exc, ResearchSynthesisError):
            raise
        raise failure from exc


async def enqueue_research_synthesis(
    *,
    kb_id: str,
    current_user: User,
    query: str,
    top_k: int,
    recall_top_k: int,
    year_from: int | None,
    year_to: int | None,
    model_spec: str | None,
    reranker_model: str | None,
    parent_run_id: str | None = None,
) -> dict[str, Any]:
    query = query.strip()
    if not query or len(query) > 4_000:
        raise ResearchSynthesisError("synthesis_invalid_query", "研究问题长度必须在 1 到 4000 个字符之间")
    if year_from is not None and year_to is not None and year_from > year_to:
        raise ResearchSynthesisError("synthesis_invalid_filter", "起始年份不能大于结束年份")
    if top_k < 2 or top_k > 20 or recall_top_k < top_k or recall_top_k > 200:
        raise ResearchSynthesisError("synthesis_invalid_retrieval_config", "top_k 或 recall_top_k 超出综述允许范围")

    await _ensure_access(current_user, kb_id)
    kb = await KnowledgeBaseRepository().get_by_kb_id(kb_id)
    if kb is None:
        raise ResearchSynthesisError("knowledge_base_not_found", "知识库不存在")
    resolved_model = (model_spec or kb.llm_model_spec or config.default_model or "").strip()
    resolved_reranker = (reranker_model or config.reranker or "").strip()
    model_info = model_cache.get_model_info(resolved_model)
    reranker_info = model_cache.get_model_info(resolved_reranker)
    if model_info is None or model_info.model_type != "chat" or not model_info.api_key:
        raise ResearchSynthesisError("synthesis_model_unavailable", "综述聊天模型未配置或缺少 API Key")
    if reranker_info is None or reranker_info.model_type != "rerank" or not reranker_info.api_key:
        raise ResearchSynthesisError("synthesis_reranker_unavailable", "综述重排模型未配置或缺少 API Key")

    uid = str(current_user.uid)
    repo = ResearchSynthesisRepository()
    active = await repo.get_active(kb_id=kb_id, uid=uid)
    if active is not None:
        raise ResearchSynthesisError("synthesis_active", f"当前知识库已有综述运行中：{active.run_id}")

    retrieval_config = {
        "mode": LOCAL_HYBRID_MODE,
        "top_k": top_k,
        "recall_top_k": recall_top_k,
        "year_from": year_from,
        "year_to": year_to,
    }
    model_config = {
        "model": resolved_model,
        "provider": model_info.provider_type,
        "reranker_model": resolved_reranker,
        "reranker_provider": reranker_info.provider_type,
    }
    run_id = uuid.uuid4().hex
    try:
        await repo.create(
            run_id=run_id,
            parent_run_id=parent_run_id,
            kb_id=kb_id,
            uid=uid,
            raw_query=query,
            model_config=model_config,
            retrieval_config=retrieval_config,
        )
    except IntegrityError as exc:
        active = await repo.get_active(kb_id=kb_id, uid=uid)
        if active is not None:
            raise ResearchSynthesisError("synthesis_active", f"当前知识库已有综述运行中：{active.run_id}") from exc
        raise ResearchSynthesisError("synthesis_record_failed", "综述运行记录创建失败") from exc

    try:
        task, _ = await tasker.enqueue_unique_by_payload(
            name=f"证据约束研究综述 ({query[:40]})",
            task_type="research_synthesis",
            payload={"run_id": run_id, "kb_id": kb_id, "uid": uid},
            payload_match={"run_id": run_id},
            statuses={"pending", "running"},
            coroutine=_resume_research_synthesis_task,
        )
        await repo.update(run_id, {"task_id": task.id})
    except Exception as exc:
        await repo.update(
            run_id,
            {
                "status": "failed",
                "stage": "failed",
                "error_type": "synthesis_enqueue_failed",
                "error_message": "综述任务提交失败",
                "completed_at": utc_now_naive(),
            },
        )
        raise ResearchSynthesisError("synthesis_enqueue_failed", "综述任务提交失败") from exc
    return {"run_id": run_id, "task_id": task.id, "status": task.status}


async def _resume_research_synthesis_task(context: TaskContext) -> dict[str, Any]:
    run_id = str(context.payload.get("run_id") or "").strip()
    if not run_id:
        raise ValueError("研究综述恢复缺少运行标识")

    repo = ResearchSynthesisRepository()
    record = await repo.get(run_id)
    if record is None:
        raise RuntimeError("研究综述运行不存在，无法恢复")
    if record.status == "success" and isinstance(record.result, dict):
        return record.result
    if record.status in {"failed", "cancelled"}:
        raise RuntimeError(f"研究综述已处于 {record.status} 状态，不能恢复")

    models = record.model_config_json or {}
    retrieval = record.retrieval_config or {}
    user = await UserRepository().get_by_uid(str(record.uid))
    if user is None or bool(user.is_deleted) or not models.get("model") or not models.get("reranker_model"):
        error = "恢复综述所需的用户或模型配置不存在"
        await repo.update(
            run_id,
            {
                "status": "failed",
                "stage": "failed",
                "error_type": "synthesis_recovery_invalid",
                "error_message": error,
                "completed_at": utc_now_naive(),
            },
        )
        raise RuntimeError(error)
    try:
        await _ensure_access(user, str(record.kb_id))
    except HTTPException as exc:
        error = "综述任务所有者已失去知识库访问权限"
        await repo.update(
            run_id,
            {
                "status": "failed",
                "stage": "failed",
                "error_type": "synthesis_recovery_invalid",
                "error_message": error,
                "completed_at": utc_now_naive(),
            },
        )
        raise RuntimeError(error) from exc

    return await _run_synthesis(
        context,
        run_id=run_id,
        kb_id=str(record.kb_id),
        current_user=user,
        query=str(record.raw_query),
        model_spec=str(models["model"]),
        reranker_model=str(models["reranker_model"]),
        retrieval_config=dict(retrieval),
        persisted_snapshot=record.retrieval_snapshot,
        resume_stage_timings=record.stage_timings,
    )


async def _authorized_run(run_id: str, current_user: User):
    record = await ResearchSynthesisRepository().get(run_id)
    if record is None:
        raise ResearchSynthesisError("synthesis_run_not_found", "研究综述运行不存在")
    await _ensure_access(current_user, str(record.kb_id))
    if str(record.uid) != str(current_user.uid) and current_user.role not in {"admin", "superadmin"}:
        raise ResearchSynthesisError("forbidden", "无权查看该研究综述运行")
    return record


async def get_research_synthesis(*, run_id: str, current_user: User) -> dict[str, Any]:
    record = await _authorized_run(run_id, current_user)
    return ResearchSynthesisRepository.serialize(record)


async def cancel_research_synthesis(*, run_id: str, current_user: User) -> dict[str, Any]:
    record = await _authorized_run(run_id, current_user)
    if record.status in {"success", "failed", "cancelled"}:
        raise ResearchSynthesisError("synthesis_not_cancellable", "该研究综述已结束，不能取消")
    if not record.task_id:
        raise ResearchSynthesisError("synthesis_task_missing", "研究综述任务尚未就绪，不能取消")
    if not await tasker.cancel_task(str(record.task_id)):
        raise ResearchSynthesisError("synthesis_not_cancellable", "该研究综述已结束，不能取消")

    cancelled = await ResearchSynthesisRepository().mark_cancelled(run_id, completed_at=utc_now_naive())
    if cancelled is None:
        terminal = await ResearchSynthesisRepository().get(run_id)
        if terminal is not None and terminal.status == "cancelled":
            return ResearchSynthesisRepository.serialize(terminal)
        raise ResearchSynthesisError("synthesis_not_cancellable", "该研究综述已结束，不能取消")
    return ResearchSynthesisRepository.serialize(cancelled)


async def list_research_syntheses(*, kb_id: str, current_user: User, offset: int, limit: int) -> dict[str, Any]:
    await _ensure_access(current_user, kb_id)
    items, total = await ResearchSynthesisRepository().list_for_user(
        kb_id=kb_id,
        uid=str(current_user.uid),
        offset=offset,
        limit=limit,
    )
    return {
        "items": [ResearchSynthesisRepository.serialize(item) for item in items],
        "total": total,
        "offset": offset,
        "limit": limit,
        "has_more": offset + len(items) < total,
    }


async def regenerate_research_synthesis(*, run_id: str, current_user: User) -> dict[str, Any]:
    record = await _authorized_run(run_id, current_user)
    if str(record.uid) != str(current_user.uid):
        raise ResearchSynthesisError("forbidden", "管理员不能代替其他用户重新生成综述")
    if record.status not in {"success", "failed", "cancelled"}:
        raise ResearchSynthesisError("synthesis_active", "当前综述尚未结束，不能重新生成")
    retrieval = record.retrieval_config or {}
    models = record.model_config_json or {}
    return await enqueue_research_synthesis(
        kb_id=str(record.kb_id),
        current_user=current_user,
        query=str(record.raw_query),
        top_k=int(retrieval.get("top_k") or 8),
        recall_top_k=int(retrieval.get("recall_top_k") or 50),
        year_from=retrieval.get("year_from"),
        year_to=retrieval.get("year_to"),
        model_spec=str(models.get("model") or ""),
        reranker_model=str(models.get("reranker_model") or ""),
        parent_run_id=str(record.run_id),
    )


def _report_references(result: dict[str, Any]) -> tuple[dict[str, int], list[dict[str, Any]]]:
    references: list[dict[str, Any]] = []
    reference_numbers: dict[str, int] = {}
    for field in ("claims", "contradictions", "limitations", "research_gaps"):
        for item in result.get(field) or []:
            for citation in item.get("evidence") or []:
                chunk_id = str(citation["chunk_id"])
                if chunk_id not in reference_numbers:
                    reference_numbers[chunk_id] = len(references) + 1
                    references.append(citation)
    return reference_numbers, references


def _citation_suffix(item: dict[str, Any], reference_numbers: dict[str, int]) -> str:
    numbers = [reference_numbers[str(citation["chunk_id"])] for citation in item.get("evidence") or []]
    return " " + "".join(f"[{number}]" for number in numbers) if numbers else ""


def _markdown_inline(value: Any) -> str:
    normalized = " ".join(str(value or "").split())
    return re.sub(r"([\\`<>\[\]()!])", r"\\\1", normalized)


def _render_markdown(result: dict[str, Any]) -> str:
    reference_numbers, references = _report_references(result)
    coverage = result.get("coverage") or {}
    lines = [
        "# 证据约束研究综述",
        "",
        f"**研究问题：** {_markdown_inline(result.get('question'))}",
        "",
        "## 执行摘要",
        "",
        _markdown_inline((result.get("executive_summary") or {}).get("text")),
        "",
        "## 核心结论",
        "",
    ]
    for claim in result.get("claims") or []:
        lines.append(
            f"- **{_markdown_inline(claim['claim_id'])} · {_markdown_inline(claim['confidence'])}** "
            f"{_markdown_inline(claim['statement'])}"
            f"{_citation_suffix(claim, reference_numbers)}"
        )
    sections = (
        ("主题综述", "themes"),
        ("证据矛盾", "contradictions"),
        ("研究局限", "limitations"),
        ("研究空白", "research_gaps"),
    )
    for title, field in sections:
        lines.extend(["", f"## {title}", ""])
        for item in result.get(field) or []:
            if field == "themes":
                lines.append(f"### {_markdown_inline(item['title'])}")
                lines.extend(["", _markdown_inline(item["summary"]), ""])
            else:
                basis = f"（依据：{_markdown_inline(item['basis'])}）" if item.get("basis") else ""
                lines.append(
                    f"- {_markdown_inline(item['statement'])}{basis}{_citation_suffix(item, reference_numbers)}"
                )
    lines.extend(["", "## 未支持结论", ""])
    for item in result.get("unsupported_claims") or []:
        lines.append(f"- {_markdown_inline(item['statement'])}（{_markdown_inline(item['reason'])}）")
    lines.extend(
        [
            "",
            "## 覆盖率",
            "",
            f"- 引用覆盖率：{float(coverage.get('citation_coverage_ratio') or 0) * 100:.1f}%",
            f"- 已验证结论：{coverage.get('supported_conclusions', 0)} / {coverage.get('total_conclusions', 0)}",
            f"- 引用论文：{coverage.get('distinct_papers', 0)} / 检索论文：{coverage.get('retrieved_papers', 0)}",
            "",
            "## 方法边界",
            "",
        ]
    )
    lines.extend(f"- {_markdown_inline(item)}" for item in result.get("methodological_constraints") or [])
    lines.extend(["", "## 证据索引", ""])
    for number, citation in enumerate(references, start=1):
        section = citation.get("section_title") or citation.get("section_type") or "论文内容"
        lines.append(
            f"[{number}] {_markdown_inline(citation['paper_title'])} · {_markdown_inline(section)} · "
            f"paper_id={_markdown_inline(citation['paper_id'])} · "
            f"chunk_id={_markdown_inline(citation['chunk_id'])}"
        )
    return "\n".join(lines).strip() + "\n"


def _render_docx(result: dict[str, Any]) -> bytes:
    reference_numbers, references = _report_references(result)
    document = Document()
    document.add_heading("证据约束研究综述", 0)
    document.add_paragraph(f"研究问题：{result.get('question') or ''}")
    document.add_heading("执行摘要", level=1)
    document.add_paragraph(str((result.get("executive_summary") or {}).get("text") or ""))
    document.add_heading("核心结论", level=1)
    for claim in result.get("claims") or []:
        document.add_paragraph(
            f"{claim['claim_id']} · {claim['confidence']}：{claim['statement']}"
            f"{_citation_suffix(claim, reference_numbers)}",
            style="List Bullet",
        )
    for title, field in (
        ("主题综述", "themes"),
        ("证据矛盾", "contradictions"),
        ("研究局限", "limitations"),
        ("研究空白", "research_gaps"),
    ):
        document.add_heading(title, level=1)
        for item in result.get(field) or []:
            if field == "themes":
                document.add_heading(item["title"], level=2)
                document.add_paragraph(item["summary"])
            else:
                basis = f"（依据：{item['basis']}）" if item.get("basis") else ""
                document.add_paragraph(
                    f"{item['statement']}{basis}{_citation_suffix(item, reference_numbers)}", style="List Bullet"
                )
    document.add_heading("未支持结论", level=1)
    for item in result.get("unsupported_claims") or []:
        document.add_paragraph(f"{item['statement']}（{item['reason']}）", style="List Bullet")
    coverage = result.get("coverage") or {}
    document.add_heading("覆盖率", level=1)
    document.add_paragraph(
        f"引用覆盖率 {float(coverage.get('citation_coverage_ratio') or 0) * 100:.1f}%；"
        f"已验证结论 {coverage.get('supported_conclusions', 0)} / {coverage.get('total_conclusions', 0)}；"
        f"引用论文 {coverage.get('distinct_papers', 0)} / 检索论文 {coverage.get('retrieved_papers', 0)}。"
    )
    document.add_heading("方法边界", level=1)
    for item in result.get("methodological_constraints") or []:
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("证据索引", level=1)
    for number, citation in enumerate(references, start=1):
        section = citation.get("section_title") or citation.get("section_type") or "论文内容"
        document.add_paragraph(
            f"[{number}] {citation['paper_title']} · {section} · "
            f"paper_id={citation['paper_id']} · chunk_id={citation['chunk_id']}"
        )
    output = BytesIO()
    document.save(output)
    return output.getvalue()


async def export_research_synthesis(*, run_id: str, current_user: User, export_format: str) -> tuple[str, bytes, str]:
    record = await _authorized_run(run_id, current_user)
    result = record.result if isinstance(record.result, dict) else None
    if record.status != "success" or not result or (result.get("validation") or {}).get("status") != "verified":
        raise ResearchSynthesisError("synthesis_not_exportable", "只有验证成功的研究综述可以导出")
    filename_base = f"research_synthesis_{record.run_id[:12]}"
    if export_format == "markdown":
        return f"{filename_base}.md", _render_markdown(result).encode("utf-8"), "text/markdown; charset=utf-8"
    if export_format == "docx":
        return (
            f"{filename_base}.docx",
            _render_docx(result),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    raise ResearchSynthesisError("synthesis_export_format_invalid", "导出格式只支持 markdown 或 docx")


async def recover_research_synthesis_runs() -> int:
    repo = ResearchSynthesisRepository()
    recovered = 0
    for record in await repo.list_recoverable():
        user = await UserRepository().get_by_uid(str(record.uid))
        models = record.model_config_json or {}
        retrieval = record.retrieval_config or {}
        if user is None or bool(user.is_deleted) or not models.get("model") or not models.get("reranker_model"):
            await repo.update(
                str(record.run_id),
                {
                    "status": "failed",
                    "stage": "failed",
                    "error_type": "synthesis_recovery_invalid",
                    "error_message": "恢复综述所需的用户或模型配置不存在",
                    "completed_at": utc_now_naive(),
                },
            )
            continue
        try:
            await _ensure_access(user, str(record.kb_id))
        except HTTPException:
            await repo.update(
                str(record.run_id),
                {
                    "status": "failed",
                    "stage": "failed",
                    "error_type": "synthesis_recovery_invalid",
                    "error_message": "综述任务所有者已失去知识库访问权限",
                    "completed_at": utc_now_naive(),
                },
            )
            continue

        async def run(
            context: TaskContext,
            item=record,
            current_user=user,
            model_spec=str(models["model"]),
            reranker_model=str(models["reranker_model"]),
            run_retrieval=dict(retrieval),
        ):
            return await _run_synthesis(
                context,
                run_id=str(item.run_id),
                kb_id=str(item.kb_id),
                current_user=current_user,
                query=str(item.raw_query),
                model_spec=model_spec,
                reranker_model=reranker_model,
                retrieval_config=run_retrieval,
                persisted_snapshot=getattr(item, "retrieval_snapshot", None),
                resume_stage_timings=getattr(item, "stage_timings", None),
            )

        try:
            task, created = await tasker.enqueue_unique_by_payload(
                name=f"恢复证据约束研究综述 ({str(record.raw_query)[:40]})",
                task_type="research_synthesis",
                payload={"run_id": record.run_id, "kb_id": record.kb_id, "uid": record.uid},
                payload_match={"run_id": record.run_id},
                statuses={"pending", "running"},
                coroutine=run,
            )
            if str(getattr(record, "task_id", None) or "") != str(task.id):
                await repo.update(str(record.run_id), {"task_id": task.id})
            if created:
                recovered += 1
        except Exception:
            await repo.update(
                str(record.run_id),
                {
                    "status": "failed",
                    "stage": "failed",
                    "error_type": "synthesis_recovery_failed",
                    "error_message": "综述恢复任务提交失败",
                    "completed_at": utc_now_naive(),
                },
            )
    return recovered


tasker.register_resumable_handler("research_synthesis", _resume_research_synthesis_task)


__all__ = [
    "ResearchSynthesisError",
    "enqueue_research_synthesis",
    "export_research_synthesis",
    "get_research_synthesis",
    "list_research_syntheses",
    "recover_research_synthesis_runs",
    "regenerate_research_synthesis",
]
